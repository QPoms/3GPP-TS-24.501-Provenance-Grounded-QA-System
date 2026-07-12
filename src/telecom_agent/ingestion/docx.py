"""Structure-aware DOCX ingestion for 3GPP specifications."""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Iterator

from telecom_agent.models import SpecChunk


SECTION_TOKEN = r"\d+[A-Z]?(?:\.[0-9A-Z]+)*"
SECTION_PATTERN = re.compile(rf"^(?P<section>{SECTION_TOKEN})\s+\t?\s*(?P<title>.+)$")
HEADING_LEVEL_PATTERN = re.compile(r"Heading\s+(?P<level>\d+)", re.IGNORECASE)


@dataclass(frozen=True)
class IngestionResult:
    source_path: str
    output_path: str
    source_sha256: str
    paragraphs: int
    tables: int
    sections: int
    chunks: int
    chunks_with_page: int

    def to_dict(self) -> dict[str, Any]:
        return {
            "source_path": self.source_path,
            "output_path": self.output_path,
            "source_sha256": self.source_sha256,
            "paragraphs": self.paragraphs,
            "tables": self.tables,
            "sections": self.sections,
            "chunks": self.chunks,
            "chunks_with_page": self.chunks_with_page,
        }


@dataclass(frozen=True)
class _Block:
    kind: str
    text: str
    style: str = ""


def _iter_blocks(document: Any) -> Iterator[_Block]:
    """Yield paragraphs and tables in document order."""

    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = " ".join(paragraph.text.split())
            if text:
                yield _Block("paragraph", text, paragraph.style.name or "")
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            rows: list[str] = []
            for row in table.rows:
                cells = [" ".join(cell.text.split()) for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                yield _Block("table", "\n".join(rows), "Table")


def _toc_pages(document: Any) -> dict[str, int]:
    pages: dict[str, int] = {}
    for paragraph in document.paragraphs:
        style = (paragraph.style.name or "").casefold()
        if style.startswith("heading"):
            break
        if not style.startswith("toc"):
            continue
        parsed = _parse_toc_line(paragraph.text)
        if parsed:
            section, page = parsed
            pages[section] = page
    return pages


def _parse_toc_line(text: str) -> tuple[str, int] | None:
    parts = [part.strip() for part in text.split("\t") if part.strip()]
    if len(parts) < 3 or not re.fullmatch(SECTION_TOKEN, parts[0]) or not parts[-1].isdigit():
        return None
    return parts[0], int(parts[-1])


def _heading(block: _Block) -> tuple[int, str, str] | None:
    level_match = HEADING_LEVEL_PATTERN.fullmatch(block.style.strip())
    if not level_match:
        return None
    text = block.text.strip()
    if text.casefold() == "foreword":
        return 1, "Foreword", "Foreword"
    match = SECTION_PATTERN.match(text)
    if not match:
        return None
    return int(level_match.group("level")), match.group("section"), match.group("title").strip()


def _heading_path(stack: list[tuple[int, str]], level: int, section: str) -> tuple[str, ...]:
    while stack and stack[-1][0] >= level:
        stack.pop()
    stack.append((level, section))
    return tuple(value for _, value in stack)


def _stable_chunk_id(spec_id: str, version: str, section: str, index: int, text: str) -> str:
    digest = hashlib.sha1(text.encode("utf-8")).hexdigest()[:10]
    normalized_spec = spec_id.casefold().replace("3gpp", "").replace("ts", "").replace(".", "")
    normalized_spec = "".join(normalized_spec.split())
    safe_section = section.replace(".", "-").casefold()
    return f"ts{normalized_spec}-v{version}-{safe_section}-{index:03d}-{digest}"


def _emit_section_chunks(
    *,
    blocks: list[_Block],
    spec_id: str,
    release: str,
    version: str,
    section: str,
    title: str,
    path: tuple[str, ...],
    page: int | None,
    source_file: str,
    max_chars: int,
) -> Iterable[SpecChunk]:
    groups: list[tuple[str, list[str]]] = []
    current_kind = "paragraph"
    current: list[str] = []
    current_size = 0

    def flush() -> None:
        nonlocal current, current_size
        if current:
            groups.append((current_kind, current))
            current = []
            current_size = 0

    for block in blocks:
        if block.kind == "table":
            flush()
            groups.append(("table", [block.text]))
            continue
        if current and current_size + len(block.text) + 2 > max_chars:
            flush()
        current_kind = "paragraph"
        current.append(block.text)
        current_size += len(block.text) + 2
    flush()

    for index, (kind, parts) in enumerate(groups, start=1):
        text = "\n\n".join(parts).strip()
        if not text:
            continue
        yield SpecChunk(
            chunk_id=_stable_chunk_id(spec_id, version, section, index, text),
            spec_id=spec_id,
            release=release,
            version=version,
            section=section,
            section_title=title,
            heading_path=path,
            page_start=page,
            page_end=page,
            content_type=kind,
            text=text,
            source_file=source_file,
        )


def ingest_docx(
    source_path: str | Path,
    output_path: str | Path,
    *,
    spec_id: str = "3GPP TS 24.501",
    release: str = "Rel-19",
    version: str = "19.2.0",
    max_chars: int = 2400,
) -> IngestionResult:
    """Parse a 3GPP DOCX into section-aware JSONL chunks."""

    try:
        from docx import Document
    except ImportError as error:  # pragma: no cover - environment-specific error path
        raise RuntimeError("DOCX ingestion requires the 'python-docx' package") from error

    source = Path(source_path)
    output = Path(output_path)
    if not source.is_file():
        raise FileNotFoundError(f"Specification document not found: {source}")
    if max_chars < 500:
        raise ValueError("max_chars must be at least 500")

    document = Document(source)
    toc_pages = _toc_pages(document)
    stack: list[tuple[int, str]] = []
    current_section = ""
    current_title = ""
    current_path: tuple[str, ...] = ()
    current_blocks: list[_Block] = []
    chunks: list[SpecChunk] = []
    paragraph_count = table_count = section_count = 0

    def flush_section() -> None:
        nonlocal current_blocks
        if not current_section:
            current_blocks = []
            return
        chunks.extend(
            _emit_section_chunks(
                blocks=current_blocks,
                spec_id=spec_id,
                release=release,
                version=version,
                section=current_section,
                title=current_title,
                path=current_path,
                page=toc_pages.get(current_section),
                source_file=source.name,
                max_chars=max_chars,
            )
        )
        current_blocks = []

    body_started = False
    for block in _iter_blocks(document):
        heading = _heading(block)
        if heading:
            body_started = True
            flush_section()
            level, current_section, current_title = heading
            current_path = _heading_path(stack, level, current_section)
            section_count += 1
            continue
        if not body_started:
            continue
        current_blocks.append(block)
        if block.kind == "table":
            table_count += 1
        else:
            paragraph_count += 1
    flush_section()

    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="\n") as handle:
        for chunk in chunks:
            handle.write(json.dumps(chunk.to_dict(), ensure_ascii=False) + "\n")

    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    return IngestionResult(
        source_path=str(source.resolve()),
        output_path=str(output.resolve()),
        source_sha256=source_hash,
        paragraphs=paragraph_count,
        tables=table_count,
        sections=section_count,
        chunks=len(chunks),
        chunks_with_page=sum(chunk.page_start is not None for chunk in chunks),
    )
