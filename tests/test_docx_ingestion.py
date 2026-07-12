from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

from telecom_agent.ingestion.docx import _parse_toc_line, ingest_docx


@unittest.skipIf(Document is None, "python-docx is not installed")
class DocxIngestionTests(unittest.TestCase):
    def test_parses_alphanumeric_3gpp_toc_section(self) -> None:
        self.assertEqual(_parse_toc_line("4.2A.1\tGeneral\t53"), ("4.2A.1", 53))

    def test_preserves_sections_and_tables(self) -> None:
        assert Document is not None
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "sample.docx"
            output = Path(temp_dir) / "chunks.jsonl"
            document = Document()
            document.add_heading("1\tScope", level=1)
            document.add_paragraph("This specification defines a sample procedure.")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Field"
            table.cell(0, 1).text = "Value"
            document.add_heading("1.1\tDetails", level=2)
            document.add_paragraph("The UE sends a sample message.")
            document.save(source)

            result = ingest_docx(source, output, max_chars=500)
            chunks = [json.loads(line) for line in output.read_text("utf-8").splitlines()]

        self.assertEqual(result.sections, 2)
        self.assertEqual(result.tables, 1)
        self.assertEqual({chunk["section"] for chunk in chunks}, {"1", "1.1"})
        self.assertIn("table", {chunk["content_type"] for chunk in chunks})


if __name__ == "__main__":
    unittest.main()
